"""Bounded local extraction for team-owned files. Not authoritative legal text without review."""
from __future__ import annotations

import csv
import mimetypes
import zipfile
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath


@dataclass(frozen=True)
class Extract:
    source_name: str
    source_hash: str
    text: str
    locator: str
    status: str = 'needs_review'

class Extractor:
    def __init__(self, max_bytes: int = 25_000_000, max_entries: int = 100,
                 max_depth: int = 2, ocr=None, transcribe=None):
        self.max_bytes,self.max_entries,self.max_depth = max_bytes,max_entries,max_depth
        self.ocr,self.transcribe = ocr,transcribe
    def read(self, path: str | Path) -> tuple[Extract,...]:
        p=Path(path)
        if p.stat().st_size > self.max_bytes: raise ValueError('file too large')
        return self._bytes(p.name,p.read_bytes(),0)
    def _bytes(self,name: str,data: bytes,depth: int) -> tuple[Extract,...]:
        if len(data)>self.max_bytes or depth>self.max_depth: raise ValueError('extraction limit')
        suffix=Path(name).suffix.lower()
        digest=sha256(data).hexdigest()
        if suffix=='.zip':
            out: list[Extract] = []; total=0
            with zipfile.ZipFile(BytesIO(data)) as archive:
                infos=archive.infolist()
                if len(infos)>self.max_entries: raise ValueError('too many archive entries')
                for item in infos:
                    rel=PurePosixPath(item.filename)
                    if (rel.is_absolute() or '..' in rel.parts or '\\' in item.filename
                        or item.file_size>self.max_bytes): raise ValueError('unsafe archive entry')
                    if item.is_dir(): continue
                    total+=item.file_size
                    if total>self.max_bytes: raise ValueError('archive expansion limit')
                    content=archive.read(item)
                    out.extend(self._bytes(name+'!'+item.filename,content,depth+1))
            return tuple(out)
        if suffix in ('.txt','.md','.json','.html','.xml'):
            return (Extract(name,digest,data.decode('utf-8'), 'text:entire', 'needs_review'),)
        if suffix=='.csv':
            lines=data.decode('utf-8-sig').splitlines()
            return tuple(Extract(name,digest,cell,f'row:{i},column:{j}')
                         for i,row in enumerate(csv.reader(lines),1)
                         for j,cell in enumerate(row,1) if cell)
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(BytesIO(data),strict=True)
            if len(reader.pages)>1000: raise ValueError('too many pages')
            return tuple(Extract(name,digest,text,f'page:{i}')
                         for i,page in enumerate(reader.pages,1)
                         if (text:=page.extract_text()))
        if suffix=='.docx':
            from docx import Document
            document=Document(BytesIO(data))
            return tuple(Extract(name,digest,p.text,f'paragraph:{i}')
                         for i,p in enumerate(document.paragraphs,1) if p.text)
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(BytesIO(data),read_only=True,data_only=True)
            if len(book.sheetnames)>100: raise ValueError('too many sheets')
            return tuple(Extract(name,digest,str(cell.value),f'{sheet.title}!{cell.coordinate}')
                         for sheet in book for row in sheet.iter_rows(max_row=10000,max_col=100)
                         for cell in row if cell.value is not None)
        mime,_=mimetypes.guess_type(name)
        if mime and mime.startswith('image/') and self.ocr:
            return (Extract(name,digest,self.ocr(data,mime),'image:ocr'),)
        if mime and (mime.startswith(('audio/', 'video/'))) and self.transcribe:
            return (Extract(name,digest,self.transcribe(data,mime),'media:transcript'),)
        raise ValueError('unsupported format or missing OCR/transcription adapter')
